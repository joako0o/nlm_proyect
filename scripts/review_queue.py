"""Seguimiento reproducible de las 783 alertas originales; no cierre masivo.

Distingue decisiones de lectura de comparaciones automáticas. Los offsets por
padre omiten sólo espacios y permiten seguir segmentos cuando cambian los IDs.
No suprime alertas del constructor ni reasigna actores.
"""
import collections
import csv
import hashlib
import json
import re
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from curation import normalize_quote, speaker_intervals, SPEAKER_REVIEW_SOURCE
from procedural import FORMULA_REVIEWS
from mention_reviews import annotation_for, FIELDS as MENTION_FIELDS

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT/'data/curation/cola_783.json'
DECISIONS = ROOT/'data/curation/revisiones_cola_783.json'


def compact(t):
    return re.sub(r'\s+', '', t)


def load_baseline(path=BASE, decisions_path=DECISIONS):
    baseline = json.loads(Path(path).read_text(encoding='utf-8'))
    rows = baseline['Filas']
    if len(rows)!=783 or len({r['ID'] for r in rows})!=783:
        raise ValueError('La instantánea debe conservar 783 IDs únicos')
    by_id = {r['ID']:r for r in rows}
    for r in rows:
        if hashlib.sha256(r['Texto'].encode()).hexdigest()!=r['SHA256_Texto']:
            raise ValueError('Instantánea: hash de texto inválido')
        if r['Fin_Compacto']-r['Inicio_Compacto']!=len(compact(r['Texto'])):
            raise ValueError('Instantánea: intervalo inválido')
    decisions = {}
    for e in json.loads(Path(decisions_path).read_text(encoding='utf-8')):
        rid = e['ID_Original']
        if rid in decisions or rid not in by_id or e['SHA256_Texto']!=by_id[rid]['SHA256_Texto']:
            raise ValueError('Decisión de cola duplicada o sin texto aplicable')
        if e['Decision'] not in {'BREVE_VALIDO_REVISADO','CONTINUIDAD_ENTRE_PADRES_REVISADA','MENCION_LEGITIMA_REVISADA'} or not e['Justificacion']:
            raise ValueError('Decisión de cola inválida')
        if e['Cita'] not in by_id[rid]['Texto']:
            raise ValueError('Cita de cola inexistente')
        decisions[rid]=e
    return baseline, decisions


def track(rows, baseline, decisions, speaker_reviews=None):
    speaker_reviews = speaker_reviews or {}
    intervals=collections.defaultdict(list)
    positions=collections.Counter()
    for i,r in enumerate(rows):
        p=r['ID_Padre'];start=positions[p];end=start+len(compact(r['Texto']));positions[p]=end
        intervals[p].append((start,end,i,r))
    result=[]
    for old in baseline['Filas']:
        start,end=old['Inicio_Compacto'],old['Fin_Compacto']
        group=[(s,e,i,r) for s,e,i,r in intervals[old['ID_Padre']] if s<end and e>start]
        pieces=[compact(r['Texto'])[max(start,s)-s:min(end,e)-s] for s,e,i,r in group]
        if ''.join(pieces)!=compact(old['Texto']):
            raise ValueError(f'Cola {old["ID"]}: no se conserva el intervalo original')
        if any(str(r['Fecha'])[:10]!=old['Fecha'] for s,e,i,r in group):
            raise ValueError('Cola: cambió la sesión')
        now=[r for s,e,i,r in group]
        flags=sorted({f for r in now for f in (r['Motivos_Revision'] or '').split(';') if f})
        modified=len(group)!=1 or group[0][0]!=start or group[0][1]!=end or now[0]['Actor_Final']!=old['Actor_Final']
        method_changed=any(r['Fuente_Actor']!=old['Fuente_Actor'] for r in now)
        form=FORMULA_REVIEWS.get(normalize_quote(old['Texto']))
        review=decisions.get(old['ID'])
        status='PENDIENTE_LECTURA_CONTEXTUAL';kind='TRIAJE_AUTOMATICO'
        reason='No adjudicado: leer el intervalo con los párrafos contiguos. La alerta no es un error confirmado.'
        action='Revisar sujeto, destinatarios, citas y continuidad; no heredar por proximidad.'
        if form and not modified and all(r['Duplicado_Formula']=='SI' for r in now):
            status='FORMULA_PROCEDIMENTAL_RECLASIFICADA';kind='LECTURA_DIRIGIDA_POR_AGENTE'
            reason=form['Revision_ID']+': '+form['Justificacion'];action='Conservar todas las repeticiones; revisar separadamente otros motivos si quedan.'
        elif review and not modified:
            status=review['Decision'];kind='LECTURA_DIRIGIDA_POR_AGENTE';reason=review['Justificacion']
            action='Conservar texto y actor; la alerta automática queda visible y la revisión está en el registro separado.'
            if status=='CONTINUIDAD_ENTRE_PADRES_REVISADA':
                following=rows[group[-1][2]+1] if group[-1][2]+1<len(rows) else {}
                if following.get('ID_Padre')!=review['Padre_Continuacion'] or following.get('ID_Turno')!=now[-1]['ID_Turno']:
                    raise ValueError('Continuación revisada perdió el turno/antecedente')
        elif modified:
            status='SEGMENTACION_O_ACTOR_MODIFICADO';kind='COMPARACION_AUTOMATICA'
            reason='Se modificó la atribución o un límite dentro del intervalo original; no se certifica pureza de todos los segmentos resultantes.'
            action='Contrastar los nuevos segmentos; mantener pendientes las alertas residuales.'
            directed = next((e for e in speaker_intervals(speaker_reviews.get(old['ID_Padre']))
                             if any(r['Fuente_Actor']==SPEAKER_REVIEW_SOURCE and r['Actor_Final']==e['Actor'] for r in now)), None)
            if directed and any(r['Fuente_Actor']==SPEAKER_REVIEW_SOURCE and r['Actor_Final']==directed['Actor'] for r in now):
                status='CORRECCION_DIRIGIDA_APLICADA';kind='LECTURA_DIRIGIDA_POR_AGENTE'
                reason=directed['Revision_ID']+': '+directed['Justificacion']+' Se adjudica el tramo documentado, no la pureza de todo el intervalo original.'
        elif method_changed:
            status='METODO_ACTUALIZADO';kind='COMPARACION_AUTOMATICA'
            reason='Cambió la evidencia reconocida sin cambiar actor/límites; no equivale a corregir un actor.'
        elif 'VARIANTE_IDENTIDAD_POR_VERIFICAR' in old['Motivos_Revision']:
            status='IDENTIDAD_PENDIENTE';action='Verificar identidad con documentación externa; no fusionar Ricaurte automáticamente.'
        elif 'DUPLICADO_NO_FORMULA' in old['Motivos_Revision']:
            status='REPETICION_SUSTANTIVA_PENDIENTE';kind='LECTURA_DIRIGIDA_POR_AGENTE'
            reason='Repite una recomendación de mantener la TPM, no sólo una fórmula de apertura o paso de palabra.'
            action='Conservar ambas recomendaciones; verificar cada sesión antes de deduplicar.'
        priority='ALTA' if any(k in old['Motivos_Revision'] for k in ['POSIBLE_OTRO','ATRIBUCION_','VARIANTE_IDENTIDAD']) else 'MEDIA'
        if status in ['FORMULA_PROCEDIMENTAL_RECLASIFICADA','BREVE_VALIDO_REVISADO','CONTINUIDAD_ENTRE_PADRES_REVISADA','MENCION_LEGITIMA_REVISADA']:
            priority='SEGUIMIENTO' if flags else 'RESUELTO_EN_ESTE_MOTIVO'
        first,last=group[0][2],group[-1][2]
        prev=rows[first-1] if first and str(rows[first-1]['Fecha'])[:10]==old['Fecha'] else {}
        nxt=rows[last+1] if last+1<len(rows) and str(rows[last+1]['Fecha'])[:10]==old['Fecha'] else {}
        result.append(dict(ID_Original=old['ID'],ID_Padre=old['ID_Padre'],Fecha=old['Fecha'],
            Actor_Original_Cola=old['Actor_Final'],Motivos_Originales=old['Motivos_Revision'],
            Estado_Seguimiento=status,Tipo_Revision=kind,Prioridad=priority,
            IDs_Actuales=';'.join(str(r['ID']) for r in now),
            Actores_Actuales=' → '.join(r['Actor_Final'] for r in now),
            Alertas_Actuales=';'.join(flags),Justificacion=reason,Siguiente_Paso=action,
            Texto_Original=old['Texto'],Final_Anterior_Actual=(prev.get('Texto') or '')[-500:],
            Inicio_Siguiente_Actual=(nxt.get('Texto') or '')[:500],SHA256_Texto=old['SHA256_Texto']))
    return result


def build_report(rows, out, speaker_reviews=None, mention_annotations=None):
    baseline,decisions=load_baseline()
    tracked=track(rows,baseline,decisions,speaker_reviews)
    current=[{**r, **annotation_for(r,mention_annotations)} for r in rows if r['Estado_Revision']=='PENDIENTE_REVISION']
    summary={'cola_original':len(tracked),'alertas_actuales':len(current),
        'menciones_actuales_documentadas':len(mention_annotations or {}),
        'estados':dict(collections.Counter(r['Estado_Seguimiento'] for r in tracked)),
        'tipo_revision':dict(collections.Counter(r['Tipo_Revision'] for r in tracked)),
        'filas_originales_con_alertas_actuales':sum(bool(r['Alertas_Actuales']) for r in tracked),
        'nota':'Estados por intervalo original; una fila puede corresponder ahora a varios segmentos. Triaje/comparación automática no son lectura humana ni cierre semántico. Las 783 no fueron revisadas exhaustivamente.',
        'sha256_base_original':baseline['SHA256_Base']}
    (out/'resumen_revision_783.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    with (out/'revision_783.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(tracked[0]));w.writeheader();w.writerows(tracked)
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Resumen'
    ws.append(['Revisión de las 783 alertas originales','Resultado'])
    ws.append(['Advertencia',summary['nota']]);ws.append(['Alertas en la versión actual',len(current)])
    for k,v in summary['estados'].items():ws.append([k,v])
    ws.append(['Lecturas de menciones actuales (no estados de las 783)',summary['menciones_actuales_documentadas']])
    for name,items,fields in [('Seguimiento_783',tracked,list(tracked[0])),
        ('Alertas_actuales',current,['ID','ID_Padre','Fecha','Actor_Final','Motivos_Revision','ID_Turno','Texto']+list(MENTION_FIELDS))]:
        sheet=wb.create_sheet(name);sheet.append(fields)
        for r in items:sheet.append([r[k] for k in fields])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(color='FFFFFF',bold=True)
        for col in sheet.columns:
            title=str(col[0].value);sheet.column_dimensions[col[0].column_letter].width=55 if any(k in title for k in ['Texto','Justificacion','Paso','Anterior','Siguiente','Motivos','Alertas','Alcance']) else 25
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height=60
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    wb.save(out/'revision_783.xlsx')
    return summary
