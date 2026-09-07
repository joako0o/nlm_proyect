# -*- coding: utf-8 -*-
"""Genera la copia final del consolidado a partir de la base de referencia."""
import datetime as dt
from pathlib import Path
import openpyxl

REPO = Path(__file__).resolve().parent.parent
SRC = REPO / 'data' / 'processed' / 'consolidado_base_referencia.xlsx'
OUT = REPO / 'data' / 'processed' / 'consolidado_base_referencia_final.xlsx'

SOURCE_COLUMNS = [
    'ID',
    'Fecha',
    'Actor_Final',
    'Rol_Final',
    'Fuente_Actor',
    'Fuente_Rol',
    'Tipo_Acta',
    'Página',
    'Texto',
    'Tema_Categoria',
    'Palabra_Clave_Categoria',
    'Duplicado_Exacto',
    'Duplicado_Formula',
]

src_wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
src_ws = src_wb['Consolidado']
src_rows = list(src_ws.iter_rows(values_only=True))
src_header = [str(c) for c in src_rows[0]]
src_index = {k: j for j, k in enumerate(src_header)}
src_data = src_rows[1:]

out_wb = openpyxl.Workbook()
out_ws = out_wb.active
out_ws.title = 'Consolidado'
out_ws.append(SOURCE_COLUMNS)

for row in src_data:
    out_ws.append([row[src_index[col]] for col in SOURCE_COLUMNS])

for row in out_ws.iter_rows(min_row=2, max_row=out_ws.max_row, min_col=2, max_col=2):
    for cell in row:
        if isinstance(cell.value, (dt.datetime, dt.date)):
            cell.number_format = 'YYYY-MM-DD'

for col in out_ws.iter_cols():
    width = min(max(max((len(str(c.value)) for c in col[:400] if c.value is not None), default=8) + 2, 10), 60)
    out_ws.column_dimensions[col[0].column_letter].width = width

out_ws.freeze_panes = 'A2'
out_wb.save(OUT)

print(f'Copiadas {len(src_data)} filas')
print('Columnas:', SOURCE_COLUMNS)
print('Guardado', OUT)
